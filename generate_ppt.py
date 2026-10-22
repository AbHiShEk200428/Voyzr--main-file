from pptx import Presentation
from pptx.util import Inches, Pt

def create_voyzr_ppt():
    prs = Presentation()

    # Helper function to add slide
    def add_slide(title, bullet_points, diagram_placeholder_text):
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = title

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        for point in bullet_points:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
        
        # Add a text box as a placeholder for the diagram
        left = Inches(1)
        top = Inches(4.5)
        width = Inches(8)
        height = Inches(2)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = f"[DIAGRAM PLACEHOLDER: {diagram_placeholder_text}]"
        p.font.bold = True
        p.font.size = Pt(18)

    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "VOYZR: THE TRIP PLANNER"
    subtitle.text = "A Microservice-Based Group Travel Application\nPresented by: ABHISHEK VS & Team"

    # 2. Introduction
    add_slide("1. Introduction", 
              ["Voyzr is an innovative web application for group trip coordination.",
               "Built on modular microservice architecture.",
               "Objective: Simplify travel planning, real-time tracking, and expense sharing."],
              "Project Concept/Overview Image")

    # 3. Block Diagram (Architecture)
    add_slide("2. Block Diagram (System Architecture)", 
              ["Three-tier architecture: Frontend, Gateway, and Backend Services.",
               "React (Frontend) communicating via API Gateway.",
               "Distributed databases: PostgreSQL, MongoDB, Redis, S3."],
              "Architecture Overview (Frontend -> Gateway -> Services -> DB)")

    # 4. Modules Diagram
    add_slide("3. Modules & Services", 
              ["Core Services: Auth, User, Trip, Location.",
               "Engagement Services: Chat, Notification, Places, Expense.",
               "Infrastructure: Kafka (Messaging), Docker/K8s (Deployment)."],
              "Service Module Breakdown (Circle/Grid view)")

    # 5. Use Case Diagram
    add_slide("4. Use Case Diagram", 
              ["Actors: Organizer (Trip Creator) and Members.",
               "Primary Use Cases: Create Trip, Track Live Location, Chat, Split Expenses.",
               "System boundaries include Google Maps and FCM integrations."],
              "Stick figures (Organizer/Member) connected to system actions")

    # 6. ER Diagram
    add_slide("5. Entity Relationship (ER) Diagram", 
              ["Entities: User (Profile), Trip (Details), Expense (Transactions).",
               "Relationships: One-to-Many for Trips-Expenses, Many-to-Many for Users-Trips.",
               "Geospatial data storage via PostGIS."],
              "Database Tables with Foreign Key relations")

    # 7. Class Diagram
    add_slide("6. Class Diagram (Trip Service)", 
              ["Focus on the 'Trip Service' business logic.",
               "Classes: TripController, TripService, TripRepository.",
               "Models: Trip (Entity), Member (Value Object)."],
              "UML Class structure with methods and properties")

    # 8. Activity Diagram
    add_slide("7. Activity Diagram", 
              ["Workflow: Start -> Login -> Create Trip -> Invite Members -> Live Navigation -> End.",
               "Handles parallel activities like real-time chatting while traveling."],
              "Flowchart showing the step-by-step user journey")

    # 9. Library & Tech Stack
    add_slide("8. Library / Tech Stack Diagram", 
              ["Backend: Java 17, Spring Boot, Spring Security, Kafka.",
               "Frontend: React, Redux, Google Maps SDK.",
               "Infra: Docker, Kubernetes, Jenkins CI/CD."],
              "Logos/Icons representing the technologies used")

    # 10. Data Flow Diagram (DFD)
    add_slide("9. Data Flow Diagram (DFD Level 1)", 
              ["User input flows through API Gateway.",
               "Location updates processed by Kafka/Redis.",
               "Final trip data persisted in PostgreSQL."],
              "Data flow arrows between Users, Services, and Storage")

    # 11. Conclusion
    add_slide("10. Conclusion & Future Scope", 
              ["Successfully integrated microservices for seamless trip coordination.",
               "Future: AI-driven destination recommendations and personalized itineraries.",
               "Scalable architecture ready for global deployment."],
              "Final Summary/Thank You Image")

    prs.save('Voyzr_Project_Presentation.pptx')
    print("PPT created successfully: Voyzr_Project_Presentation.pptx")

if __name__ == "__main__":
    create_voyzr_ppt()
